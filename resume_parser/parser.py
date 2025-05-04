import PyPDF2
import docx
import time
import re
import os
import logging

# Configure logging
logging.basicConfig(level=logging.INFO, 
                    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

def parse_resume(file_path, use_llm=True):
    """Parse a resume file and extract relevant information."""
    start_time = time.time()
    
    # Import the model here to avoid circular imports
    from resume_parser.model import ResumeParserModel
    
    # Extract text from the file
    text = extract_text(file_path)
    
    # Clean up the text
    text = clean_text(text)
    
    # Initialize the model and parse the resume
    model = ResumeParserModel()
    parsed_data = model.parse_resume(text, use_llm)
    
    # Add degree-specific information for filtering
    parsed_data["degree_education"] = model.get_degree_education(parsed_data)
    
    # Extract degree-specific GPA and graduation year
    degree_gpa = None
    degree_graduation_year = None
    
    if parsed_data["degree_education"]:
        degree_gpa = parsed_data["degree_education"].get("gpa")
        degree_graduation_year = parsed_data["degree_education"].get("graduation_year")
    
    # Add these as separate fields for easier access
    parsed_data["degree_gpa"] = degree_gpa
    parsed_data["degree_graduation_year"] = degree_graduation_year
    
    # Log the parsing time
    parsing_time = time.time() - start_time
    print(f"Total parsing completed in {parsing_time:.2f} seconds.")
    
    return parsed_data

def extract_text(file_path):
    """Extract text from different file formats."""
    if not os.path.exists(file_path):
        logger.error(f"File not found: {file_path}")
        return ""
        
    if file_path.lower().endswith('.pdf'):
        return extract_text_from_pdf(file_path)
    elif file_path.lower().endswith('.docx'):
        return extract_text_from_docx(file_path)
    elif file_path.lower().endswith(('.png', '.jpg', '.jpeg')):
        return extract_text_from_image(file_path)
    else:
        # Try to read as plain text
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except Exception as e:
            logger.error(f"Error reading file: {e}")
            return ""

def extract_text_from_pdf(file_path):
    """Extract text from PDF file."""
    text = ""
    try:
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfReader(file)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        logger.info(f"Successfully extracted text from PDF: {file_path}")
    except Exception as e:
        logger.error(f"Error extracting text from PDF: {e}")
    return text.strip()

def extract_text_from_docx(file_path):
    """Extract text from DOCX file."""
    text = ""
    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            text += para.text + "\n"
        logger.info(f"Successfully extracted text from DOCX: {file_path}")
    except Exception as e:
        logger.error(f"Error extracting text from DOCX: {e}")
    return text.strip()

def extract_text_from_image(file_path):
    """Extract text from image file using OCR."""
    text = ""
    try:
        import pytesseract
        from PIL import Image
        
        img = Image.open(file_path)
        text = pytesseract.image_to_string(img)
        logger.info(f"Successfully extracted text from image: {file_path}")
    except ImportError:
        logger.error("pytesseract or PIL not installed. Install with: pip install pytesseract pillow")
    except Exception as e:
        logger.error(f"Error extracting text from image: {e}")
    return text.strip()

def clean_text(text):
    """Clean up the extracted text."""
    # Replace CID placeholders that often appear in PDFs
    text = re.sub(r'\(cid:[0-9]+\)', '', text)
    
    # Normalize section headers with markdown-style headers
    text = re.sub(r'\n([A-Z][A-Z\s]+)(?:\n|:)', r'\n## \1\n', text)
    
    # Ensure consistent bullet points
    text = re.sub(r'[\*\+\-]\s', '• ', text)
    
    # Replace multiple spaces with a single space
    text = re.sub(r' +', ' ', text)
    
    # Ensure consistent newlines
    text = re.sub(r'\n+', '\n', text)
    
    # Additional cleaning for better parsing
    # Remove page numbers
    text = re.sub(r'\n\s*\d+\s*\n', '\n', text)
    
    # Remove header/footer content that often appears in resumes
    text = re.sub(r'^\s*Page \d+ of \d+\s*$', '', text, flags=re.MULTILINE)
    
    return text

def get_all_skills():
    """Return a list of common skills to use for filtering."""
    return [
        # Programming Languages
        "Python", "Java", "JavaScript", "TypeScript", "C++", "C#", "Ruby", "PHP", 
        "Swift", "Kotlin", "Go", "Rust", "Scala", "R", "MATLAB", "Perl", "Shell",
        
        # Web Development
        "HTML", "CSS", "React", "Angular", "Vue.js", "Node.js", "Express", 
        "Django", "Flask", "Ruby on Rails", "Spring Boot", "ASP.NET", 
        "jQuery", "Bootstrap", "Tailwind CSS", "Material UI", "Redux",
        
        # Database & Data
        "SQL", "MySQL", "PostgreSQL", "MongoDB", "SQLite", "Oracle", 
        "Redis", "Elasticsearch", "Cassandra", "Firebase", "DynamoDB",
        "Data Analysis", "Data Science", "Data Visualization", "Power BI", "Tableau",
        
        # DevOps & Cloud
        "AWS", "Azure", "Google Cloud", "Docker", "Kubernetes", "Jenkins", 
        "CI/CD", "Git", "GitHub", "GitLab", "Terraform", "Ansible", "Puppet",
        
        # AI & Machine Learning
        "Machine Learning", "Deep Learning", "AI", "TensorFlow", "PyTorch", 
        "Keras", "scikit-learn", "NLP", "Computer Vision", "Generative AI", 
        "Prompt Engineering", "LLMs", "Neural Networks",
        
        # Mobile Development
        "iOS", "Android", "React Native", "Flutter", "Xamarin", "SwiftUI",
        
        # Soft Skills
        "Team Collaboration", "Problem-Solving", "Critical Thinking", 
        "Time Management", "Project Management", "Communication", 
        "Leadership", "Adaptability", "Creativity", "Analytical Skills",
        "Presentation Skills", "Detail-Oriented", "Strategic Thinking",
        
        # Methodologies
        "Agile", "Scrum", "Kanban", "Waterfall", "Test-Driven Development",
        "Object-Oriented Programming", "Functional Programming",
        
        # APIs & Services
        "RESTful API", "GraphQL", "SOAP", "Microservices", "Serverless",
        
        # Testing
        "Unit Testing", "Integration Testing", "Selenium", "Jest", "Mocha",
        "Cypress", "JUnit", "PyTest"
    ]

def format_parsed_data(parsed_data):
    """Format the parsed data for better display in the web interface."""
    formatted_data = {
        "skills": [],
        "education": [],
        "experience": [],
        "degree_info": None  # Add degree-specific information for filtering
    }
    
    # Format skills
    if parsed_data.get("skills"):
        all_skills = get_all_skills()
        # Prioritize skills that match our predefined list
        matched_skills = [skill for skill in parsed_data["skills"] if skill in all_skills]
        other_skills = [skill for skill in parsed_data["skills"] if skill not in all_skills]
        
        formatted_data["skills"] = matched_skills + other_skills
    
    # Format education
    if parsed_data.get("education"):
        # Separate education entries by level
        degree_education = []
        secondary_education = []
        high_school_education = []
        
        for edu in parsed_data.get("education", []):
            education_level = edu.get("education_level")
            
            if education_level == "degree":
                degree_education.append(edu)
            elif education_level == "secondary":
                secondary_education.append(edu)
            elif education_level == "high_school":
                high_school_education.append(edu)
            else:
                # If education_level is not set, try to determine from degree name
                if "degree" in edu:
                    degree_name = edu.get("degree", "").lower()
                    if any(keyword in degree_name for keyword in ['bachelor', 'b.tech', 'btech', 'engineering']):
                        degree_education.append(edu)
                    elif any(keyword in degree_name for keyword in ['12th', 'intermediate', 'senior secondary']):
                        secondary_education.append(edu)
                    elif any(keyword in degree_name for keyword in ['10th', 'high school', 'secondary']):
                        high_school_education.append(edu)
        
        # Order education by level: degree first, then secondary, then high school
        formatted_data["education"] = degree_education + secondary_education + high_school_education
        
        # Extract degree-specific education information for filtering
        if parsed_data.get("degree_education"):
            formatted_data["degree_info"] = {
                "graduation_year": parsed_data["degree_education"].get("graduation_year"),
                "degree_gpa": parsed_data["degree_education"].get("gpa")
            }
        # Fallback: use the first degree-level education if degree_education is not set
        elif degree_education:
            formatted_data["degree_info"] = {
                "graduation_year": degree_education[0].get("graduation_year"),
                "degree_gpa": degree_education[0].get("gpa")
            }
    
    # Format experience
    if parsed_data.get("experience"):
        formatted_data["experience"] = parsed_data["experience"]
    
    return formatted_data

def filter_resumes_by_criteria(resumes, skills=None, graduation_year=None, min_gpa=0):
    """Filter resumes based on specific criteria.
    
    Args:
        resumes: List of parsed resume data
        skills: List of required skills
        graduation_year: Required graduation year for degree/B.Tech
        min_gpa: Minimum GPA required for degree/B.Tech
        
    Returns:
        List of filtered resumes
    """
    filtered_resumes = []
    
    for resume in resumes:
        # Skip if no parsed data
        if not resume or not isinstance(resume, dict):
            continue
            
        # Check skills criteria
        if skills and len(skills) > 0:
            resume_skills = resume.get("skills", [])
            if not all(skill in resume_skills for skill in skills):
                continue
        
        # Get degree education information - directly from degree fields if available
        degree_graduation_year = resume.get("degree_graduation_year")
        degree_gpa = resume.get("degree_gpa")
        
        # Fallback to degree_info if direct fields not available
        if degree_graduation_year is None or degree_gpa is None:
            degree_info = resume.get("degree_info", {})
            if not degree_info:
                # Try to get from degree_education directly
                degree_education = resume.get("degree_education", {})
                if degree_education:
                    degree_graduation_year = degree_education.get("graduation_year")
                    degree_gpa = degree_education.get("gpa")
            else:
                degree_graduation_year = degree_info.get("graduation_year")
                degree_gpa = degree_info.get("degree_gpa")
        
        # Check graduation year criteria (only for degree/B.Tech)
        if graduation_year and degree_graduation_year:
            if str(degree_graduation_year) != str(graduation_year):
                continue
        
        # Check minimum GPA criteria (only for degree/B.Tech)
        if min_gpa > 0 and degree_gpa:
            try:
                degree_gpa_float = float(degree_gpa)
                if degree_gpa_float < float(min_gpa):
                    continue
            except (ValueError, TypeError):
                # If GPA can't be converted to float, it doesn't match
                continue
        
        # If all criteria pass, add to filtered results
        filtered_resumes.append(resume)
    
    return filtered_resumes

def get_degree_graduation_years(parsed_resumes):
    """Extract only graduation years from degree/B.Tech level education across all resumes."""
    graduation_years = set()
    
    for resume in parsed_resumes:
        # Try direct field first
        if resume.get("degree_graduation_year"):
            graduation_years.add(str(resume["degree_graduation_year"]))
            continue
            
        # Try degree_info next
        degree_info = resume.get("degree_info", {})
        if degree_info and degree_info.get("graduation_year"):
            graduation_years.add(str(degree_info["graduation_year"]))
            continue
            
        # Finally try degree_education
        degree_education = resume.get("degree_education", {})
        if degree_education and degree_education.get("graduation_year"):
            graduation_years.add(str(degree_education["graduation_year"]))
    
    return sorted(list(graduation_years), reverse=True)
